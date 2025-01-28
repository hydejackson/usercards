import docx as dx

# Write data to Word file
# in:  string filename (to edit), list details[[name, id, password],...]
# out: null
def putData(filename, details: list):
    # Read in the Word document
    document = dx.Document(filename)
    # List of tables in the document (basically the pages for our case)
    docTable = document.tables
    counterVar = 0
    tempTextList = []

    # Make a list of strings of text to add to each cell
    for d in range(len(details)):
        tempText = \
        "   Name: " + str(details[d][0]) + "\n\n" + \
        "   Username: " + str(details[d][1]) + "\n\n" + \
        "   Password: " + str(details[d][2])

        tempTextList.append(tempText)

    print(tempTextList)
    print(len(tempTextList))

    # Cycle through each row/column coord and change text to an entry in the tempTextList
    for d in docTable:
        for i in range(len(d.rows)):
            for j in range(len(d.columns)):
                if counterVar < len(tempTextList):
                    print(counterVar)
                    print(tempTextList[counterVar])
                    d.cell(i, j).text = tempTextList[counterVar]
                    counterVar += 1
                else:
                    break

    # Save the updated document
    document.save(filename)